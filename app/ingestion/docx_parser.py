"""DOCX extraction with python-docx.

Two things make DOCX different from PDF:

**There are no pages.** Word paginates at render time using the printer
and font metrics; the file stores no page boundaries. Rather than
pretend otherwise, pages are *synthesised* — a new one starts at an
explicit page break, or when a running character budget is exceeded —
and every resulting page is flagged ``synthetic=True`` so that a later
citation can be honest about it. The alternative, one giant page 1,
would make every citation in a 60-page agreement read "page 1".

**Structure is declared, not inferred.** Word carries real heading
styles, so section detection starts from ``style.name`` and only falls
back to the textual heuristics used for PDFs. That makes DOCX section
labels considerably more reliable than PDF ones.

Paragraphs and tables are walked in true document order by iterating the
body XML, because ``document.paragraphs`` and ``document.tables`` are
separate sequences that lose the interleaving between them.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterator, List, Optional, Tuple

from app.config import Settings
from app.ingestion.base import DocumentParser, ParseOutcome, register
from app.ingestion.errors import CorruptDocumentError, ParserUnavailableError
from app.ingestion.sections import detect_section, normalize_section
from app.logging_config import get_logger
from app.models.document import (
    BlockKind,
    DocumentMetadata,
    ExtractionMethod,
    Page,
    TextBlock,
)

logger = get_logger(__name__)

#: Characters per synthetic page. Roughly a dense A4 page of legal text;
#: the exact value matters less than being stable and documented.
SYNTHETIC_PAGE_CHARS = 3000


def _import_docx():
    try:
        import docx

        return docx
    except ImportError as exc:  # pragma: no cover
        raise ParserUnavailableError(
            "python-docx is required to read DOCX files"
        ) from exc


def _iter_body(document) -> Iterator[Tuple[str, object]]:
    """Yield ``('paragraph'|'table', item)`` in true document order."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            yield "paragraph", Paragraph(child, document)
        elif tag == "tbl":
            yield "table", Table(child, document)


def _has_page_break(paragraph) -> bool:
    """Does this paragraph carry an explicit page break?"""
    try:
        # python-docx exposes no public API for break elements.
        xml = paragraph._p.xml  # noqa: SLF001
    except Exception:  # pragma: no cover
        return False
    return 'w:br' in xml and 'type="page"' in xml


def _table_text(table) -> str:
    """Flatten a table to pipe-separated rows.

    Structure is lost; the content is not. Good enough for retrieval,
    and noted as a limitation rather than quietly dropped.
    """
    rows: List[str] = []
    for row in table.rows:
        cells = [" ".join(cell.text.split()) for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _style_heading(paragraph) -> Optional[str]:
    """A heading label from Word's own style, if the paragraph has one."""
    try:
        style = (paragraph.style.name or "").lower()
    except Exception:  # pragma: no cover
        return None
    if style.startswith("heading") or style in ("title", "subtitle"):
        text = " ".join(paragraph.text.split())
        return normalize_section(text) if text else None
    return None


def _is_list_paragraph(paragraph) -> bool:
    try:
        return "list" in (paragraph.style.name or "").lower()
    except Exception:  # pragma: no cover
        return False


@register
class DOCXParser(DocumentParser):
    extensions = (".docx",)
    name = "python-docx"

    def parse(
        self,
        path: Path,
        metadata: DocumentMetadata,
        settings: Settings,
        force_ocr: bool = False,
    ) -> ParseOutcome:
        docx = _import_docx()
        warnings: List[str] = []

        if force_ocr:
            # DOCX has no page raster to OCR; say so rather than silently
            # ignoring the flag.
            warnings.append("force_ocr was ignored: DOCX has no scanned pages")

        try:
            document = docx.Document(str(path))
        except Exception as exc:
            raise CorruptDocumentError(
                f"The DOCX file could not be opened: {exc}"
            ) from exc

        pages: List[Page] = []
        blocks: List[TextBlock] = []
        page_number = 1
        block_index = 0
        page_chars = 0
        section: Optional[str] = None

        def flush(force: bool = False) -> None:
            """Close the current synthetic page."""
            nonlocal blocks, page_number, block_index, page_chars
            if not blocks and not force:
                return
            text = "\n\n".join(b.text for b in blocks)
            pages.append(
                Page(
                    page_number=page_number,
                    blocks=blocks,
                    text=text,
                    raw_text=text if settings.keep_raw_text else None,
                    method=(
                        ExtractionMethod.NATIVE if text.strip()
                        else ExtractionMethod.NONE
                    ),
                    synthetic=True,
                )
            )
            page_number += 1
            blocks = []
            block_index = 0
            page_chars = 0

        try:
            for kind, item in _iter_body(document):
                if kind == "paragraph":
                    if _has_page_break(item):
                        flush()

                    text = " ".join(item.text.split())
                    if not text:
                        continue

                    heading = _style_heading(item)
                    if heading is None and settings.detect_sections:
                        heading = detect_section(text)

                    if heading:
                        section = heading
                        block_kind = BlockKind.HEADING
                    elif _is_list_paragraph(item):
                        block_kind = BlockKind.LIST_ITEM
                    else:
                        block_kind = BlockKind.PARAGRAPH
                else:
                    text = _table_text(item)
                    if not text:
                        continue
                    block_kind = BlockKind.TABLE

                blocks.append(
                    TextBlock(
                        text=text,
                        page_number=page_number,
                        block_index=block_index,
                        kind=block_kind,
                        section=section,
                    )
                )
                block_index += 1
                page_chars += len(text)

                if page_chars >= SYNTHETIC_PAGE_CHARS:
                    flush()
        except CorruptDocumentError:
            raise
        except Exception as exc:
            raise CorruptDocumentError(
                f"The DOCX file could not be read: {exc}"
            ) from exc

        flush()

        if not pages:
            # A structurally valid but contentless file. Emit one empty
            # page so the document shape stays consistent; the service
            # decides whether an empty document is an error.
            pages.append(
                Page(page_number=1, method=ExtractionMethod.NONE, synthetic=True)
            )
            warnings.append("The document contains no text")

        return ParseOutcome(pages=pages, parser=self.name, warnings=warnings)


__all__ = ["DOCXParser", "SYNTHETIC_PAGE_CHARS"]
